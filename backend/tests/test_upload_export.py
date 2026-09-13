"""Backend tests for Divine Leadership Press upload + export endpoints."""
import io
import os
import time
import pytest
import requests
from docx import Document as DocxDocument
from dotenv import dotenv_values

BASE_URL = (
    os.environ.get("REACT_APP_BACKEND_URL")
    or dotenv_values("/app/frontend/.env").get("REACT_APP_BACKEND_URL")
)
assert BASE_URL, "REACT_APP_BACKEND_URL is required"
BASE_URL = BASE_URL.rstrip("/")
API = f"{BASE_URL}/api"


# --------- Fixtures ---------
@pytest.fixture(scope="module")
def session():
    s = requests.Session()
    return s


@pytest.fixture(scope="module")
def user_a(session):
    ts = int(time.time() * 1000)
    email = f"TEST_user_a_{ts}@example.com"
    r = session.post(f"{API}/auth/register", json={
        "email": email, "name": "Test A", "password": "Passw0rd!"
    })
    assert r.status_code == 200, r.text
    data = r.json()
    return {"email": email, "token": data["token"], "id": data["user"]["id"]}


@pytest.fixture(scope="module")
def user_b(session):
    ts = int(time.time() * 1000) + 1
    email = f"TEST_user_b_{ts}@example.com"
    r = session.post(f"{API}/auth/register", json={
        "email": email, "name": "Test B", "password": "Passw0rd!"
    })
    assert r.status_code == 200
    data = r.json()
    return {"email": email, "token": data["token"], "id": data["user"]["id"]}


def auth(token):
    return {"Authorization": f"Bearer {token}"}


def make_docx_bytes():
    """Build an in-memory .docx with mixed formatting."""
    d = DocxDocument()
    d.add_heading("Chapter One", level=1)
    p = d.add_paragraph()
    p.add_run("Hello ").bold = True
    p.add_run("world").italic = True
    d.add_paragraph("Plain paragraph text.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# --------- Upload tests ---------
class TestUpload:
    def test_upload_requires_auth(self, session):
        files = {"file": ("a.txt", b"hi", "text/plain")}
        r = session.post(f"{API}/documents/upload", files=files)
        assert r.status_code == 401

    def test_upload_docx_preserves_html(self, session, user_a):
        files = {"file": ("test.docx", make_docx_bytes(),
                          "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        r = session.post(f"{API}/documents/upload", files=files, headers=auth(user_a["token"]))
        assert r.status_code == 200, r.text
        data = r.json()
        c = data["content"]
        assert "<h1>" in c
        assert "<strong>" in c
        assert "<em>" in c
        assert "<p>" in c
        assert data["title"] == "test"
        assert data["original_filename"] == "test.docx"
        user_a["doc_id"] = data["id"]

    def test_upload_txt_wraps_in_p(self, session, user_a):
        files = {"file": ("note.txt", b"First paragraph.\n\nSecond paragraph.", "text/plain")}
        r = session.post(f"{API}/documents/upload", files=files, headers=auth(user_a["token"]))
        assert r.status_code == 200
        c = r.json()["content"]
        assert c.startswith("<p>")
        assert c.count("<p>") >= 2

    def test_upload_pages_returns_400_with_message(self, session, user_a):
        files = {"file": ("doc.pages", b"fakepagesdata", "application/octet-stream")}
        r = session.post(f"{API}/documents/upload", files=files, headers=auth(user_a["token"]))
        assert r.status_code == 400
        detail = r.json().get("detail", "")
        assert ".docx" in detail.lower() or "pages" in detail.lower()
        assert "export" in detail.lower()

    def test_upload_pdf_returns_400(self, session, user_a):
        files = {"file": ("x.pdf", b"%PDF-1.4 fake", "application/pdf")}
        r = session.post(f"{API}/documents/upload", files=files, headers=auth(user_a["token"]))
        assert r.status_code == 400
        assert "unsupported" in r.json()["detail"].lower() or ".docx" in r.json()["detail"].lower()


# --------- Export formats listing ---------
class TestFormats:
    def test_list_formats(self, session):
        r = session.get(f"{API}/export/formats")
        assert r.status_code == 200
        data = r.json()
        keys = {t["key"] for t in data["print_trim_sizes"]}
        for required in ["5x8", "5.25x8", "5.5x8.5", "6x9", "7x10", "8.5x11"]:
            assert required in keys, f"Missing trim {required}"
        assert any(d["key"] == "epub" for d in data["digital"])


# --------- Export tests ---------
class TestExport:
    @pytest.fixture(scope="class")
    def doc_id(self, session, user_a):
        # Use uploaded docx doc if exists, else create one
        if "doc_id" in user_a:
            return user_a["doc_id"]
        files = {"file": ("exp.docx", make_docx_bytes(),
                          "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        r = session.post(f"{API}/documents/upload", files=files, headers=auth(user_a["token"]))
        assert r.status_code == 200
        return r.json()["id"]

    def test_export_requires_auth(self, session, doc_id):
        r = session.post(f"{API}/documents/{doc_id}/export?format=pdf&trim=6x9")
        assert r.status_code == 401

    @pytest.mark.parametrize("trim", ["6x9", "8.5x11", "5x8"])
    def test_export_pdf_trims(self, session, user_a, doc_id, trim):
        r = session.post(f"{API}/documents/{doc_id}/export?format=pdf&trim={trim}", headers=auth(user_a["token"]))
        assert r.status_code == 200, r.text
        assert r.headers["content-type"].startswith("application/pdf")
        assert r.content[:4] == b"%PDF"
        cd = r.headers.get("content-disposition", "")
        assert "attachment" in cd.lower()
        assert trim in cd

    def test_export_pdf_no_trim_defaults(self, session, user_a, doc_id):
        r = session.post(f"{API}/documents/{doc_id}/export?format=pdf", headers=auth(user_a["token"]))
        assert r.status_code == 200
        assert r.content[:4] == b"%PDF"

    def test_export_epub(self, session, user_a, doc_id):
        r = session.post(f"{API}/documents/{doc_id}/export?format=epub", headers=auth(user_a["token"]))
        assert r.status_code == 200, r.text
        assert r.headers["content-type"].startswith("application/epub+zip")
        assert r.content[:2] == b"PK"

    def test_export_unsupported_format(self, session, user_a, doc_id):
        r = session.post(f"{API}/documents/{doc_id}/export?format=docx", headers=auth(user_a["token"]))
        assert r.status_code == 400

    def test_export_other_user_returns_404(self, session, user_b, doc_id):
        r = session.post(f"{API}/documents/{doc_id}/export?format=pdf&trim=6x9", headers=auth(user_b["token"]))
        assert r.status_code == 404
