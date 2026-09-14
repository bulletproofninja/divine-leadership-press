"""
Iteration 17 end-to-end audit.

Covers everything the review_request calls out that isn't already covered by
/app/backend/tests/test_billing.py:
- Auth (register, login, /me, /me/referrals, /me/badge, ref-code linkage)
- Documents CRUD
- File upload -> HTML
- Export (PDF x2 trim sizes, EPUB) — verifies file headers/magic bytes
- Existing AI editor endpoints (tighten / clarity / blurb / copyedit / tools)
- Writing agent (voices, quota, chat + history persistence + clear, command)
- Agent quota decrement 5->0 + 402 on 6th call, unlimited when subscribed
- Per-chapter audiobook preview + validation errors
- Affiliate/referral leaderboard + super-admin PUT gate
- Billing plans + diagnostics + LIVE checkout session + status auth + portal
- Stripe Connect 503 "not enabled" + connected:false + auto-payout 400/failed
- Super-admin gates on /api/admin/*
"""
import os
import io
import re
import time
import uuid
import zipfile

import pytest
import requests

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "https://editorial-studio-19.preview.emergentagent.com").rstrip("/")
API = f"{BASE_URL}/api"

ADMIN_EMAIL = os.environ.get("E2E_ADMIN_EMAIL")
ADMIN_PASS = os.environ.get("E2E_ADMIN_PASSWORD")


# ---------- helpers ----------

def _u(prefix="user"):
    return f"test_iter17_{prefix}_{uuid.uuid4().hex[:8]}@example.com"


def _register(email=None, name="Iter17 Author", password="pw12345!", ref=None):
    payload = {"email": email or _u(), "name": name, "password": password}
    if ref:
        payload["referral_code"] = ref
    r = requests.post(f"{API}/auth/register", json=payload, timeout=30)
    r.raise_for_status()
    return r.json()  # {token, user}


def _login(email, password):
    r = requests.post(f"{API}/auth/login", json={"email": email, "password": password}, timeout=30)
    r.raise_for_status()
    return r.json()


def _headers(token):
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture(scope="session")
def admin_token():
    if not ADMIN_EMAIL or not ADMIN_PASS:
        pytest.skip("Privileged E2E credentials are intentionally not stored in source control")
    tok = _login(ADMIN_EMAIL, ADMIN_PASS)["token"]
    return tok


@pytest.fixture(scope="session")
def user():
    """A fresh regular user, session-scoped, with one document created."""
    reg = _register()
    tok = reg["token"]
    uid = reg["user"]["id"]
    # create one doc
    doc = requests.post(
        f"{API}/documents",
        json={"title": "TEST_iter17_doc", "content": "<h1>Chapter One</h1><p>The bell rang out at dawn across the valley.</p><h1>Chapter Two</h1><p>Elena knelt on the frost-hardened path.</p>", "format": "6x9"},
        headers=_headers(tok), timeout=30,
    ).json()
    return {"token": tok, "email": reg["user"]["email"], "id": uid, "ref": reg["user"]["referral_code"], "doc_id": doc["id"]}


# ---------- AUTH ----------

class TestAuth:
    def test_register_login_me(self):
        email = _u("auth")
        reg = _register(email=email)
        assert reg["user"]["email"] == email
        assert reg["user"]["referral_code"]
        # login
        li = _login(email, "pw12345!")
        assert li["user"]["id"] == reg["user"]["id"]
        # /me
        me = requests.get(f"{API}/auth/me", headers=_headers(li["token"]), timeout=15).json()
        assert me["email"] == email
        assert me["is_super_admin"] is False
        assert isinstance(me["referral_code"], str)

    def test_admin_me_flags(self, admin_token):
        me = requests.get(f"{API}/auth/me", headers=_headers(admin_token), timeout=15).json()
        assert me["email"] == ADMIN_EMAIL
        assert me["is_super_admin"] is True

    def test_referrals_and_badge(self, user):
        r = requests.get(f"{API}/auth/me/referrals", headers=_headers(user["token"]), timeout=15)
        assert r.status_code == 200
        assert r.json()["referral_code"] == user["ref"]

        b = requests.get(f"{API}/auth/me/badge", headers=_headers(user["token"]), timeout=15)
        assert b.status_code == 200
        assert "badge" in b.json()

    def test_referred_by_set_on_registration(self, user):
        # register a new user with the fixture user's referral code
        email = _u("ref")
        reg = _register(email=email, ref=user["ref"])
        # login and read referrals of the referrer
        r = requests.get(f"{API}/auth/me/referrals", headers=_headers(user["token"]), timeout=15).json()
        assert r["total_referred"] >= 1


# ---------- DOCUMENTS CRUD + upload + export ----------

class TestDocuments:
    def test_crud(self, user):
        h = _headers(user["token"])
        # create
        r = requests.post(f"{API}/documents", json={"title": "TEST_iter17_crud", "content": "<p>hi</p>"}, headers=h, timeout=15)
        assert r.status_code == 200
        doc = r.json()
        did = doc["id"]
        # get
        g = requests.get(f"{API}/documents/{did}", headers=h, timeout=15)
        assert g.status_code == 200
        assert g.json()["title"] == "TEST_iter17_crud"
        # list
        lst = requests.get(f"{API}/documents", headers=h, timeout=15).json()
        assert any(d["id"] == did for d in lst)
        # update
        u = requests.put(f"{API}/documents/{did}", json={"title": "TEST_iter17_crud2", "content": "<p>bye</p>"}, headers=h, timeout=15)
        assert u.status_code == 200 and u.json()["title"] == "TEST_iter17_crud2"
        # delete
        d = requests.delete(f"{API}/documents/{did}", headers=h, timeout=15)
        assert d.status_code == 200
        # verify gone
        g2 = requests.get(f"{API}/documents/{did}", headers=h, timeout=15)
        assert g2.status_code == 404

    def test_docx_upload(self, user):
        # Build a minimal valid .docx in memory
        try:
            from docx import Document as DocxDocument
        except Exception:
            pytest.skip("python-docx not available")
        buf = io.BytesIO()
        d = DocxDocument()
        d.add_heading("Uploaded Title", 0)
        d.add_paragraph("Uploaded body text.")
        d.save(buf)
        buf.seek(0)
        r = requests.post(
            f"{API}/documents/upload",
            files={"file": ("upload.docx", buf.getvalue(),
                            "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            headers=_headers(user["token"]),
            timeout=60,
        )
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["id"]
        assert "Uploaded" in (body.get("content") or "")

    def test_export_pdf_and_epub(self, user):
        h = _headers(user["token"])
        did = user["doc_id"]
        # PDF 6x9
        r = requests.post(f"{API}/documents/{did}/export?format=pdf&trim=6x9", headers=h, timeout=60)
        assert r.status_code == 200
        assert r.headers.get("content-type", "").startswith("application/pdf")
        assert r.content[:4] == b"%PDF"
        # PDF 8.5x11
        r2 = requests.post(f"{API}/documents/{did}/export?format=pdf&trim=8.5x11", headers=h, timeout=60)
        assert r2.status_code == 200
        assert r2.content[:4] == b"%PDF"
        assert r2.headers.get("x-trim-size") == "8.5x11"
        # EPUB
        r3 = requests.post(f"{API}/documents/{did}/export?format=epub", headers=h, timeout=60)
        assert r3.status_code == 200
        assert r3.headers.get("content-type", "").startswith("application/epub")
        # EPUB is a ZIP starting with "PK"
        assert r3.content[:2] == b"PK"


# ---------- AI EDITOR (existing endpoints) ----------

class TestAIEditor:
    def test_tools_list(self, user):
        r = requests.get(f"{API}/ai/tools", headers=_headers(user["token"]), timeout=15)
        assert r.status_code == 200
        keys = {t["key"] for t in r.json()["tools"]}
        assert {"tighten", "blurb"}.issubset(keys)

    def test_tighten_clarity_blurb(self, user):
        h = _headers(user["token"])
        did = user["doc_id"]
        for tool in ("tighten", "clarity", "blurb"):
            r = requests.post(f"{API}/documents/{did}/ai", json={"tool": tool}, headers=h, timeout=120)
            assert r.status_code == 200, f"{tool}: {r.text[:200]}"
            assert r.json()["result"]

    def test_copyedit(self, user):
        h = _headers(user["token"])
        r = requests.post(
            f"{API}/documents/{user['doc_id']}/copyedit",
            json={"style_guide": "chicago"},
            headers=h,
            timeout=180,
        )
        assert r.status_code == 200, r.text[:200]
        assert "issues" in r.json() and "readability" in r.json()


# ---------- WRITING AGENT ----------

class TestWritingAgent:
    def test_voices_list_has_seven(self):
        r = requests.get(f"{API}/ai/agent/voices", timeout=15)
        assert r.status_code == 200
        voices = r.json()["voices"]
        assert len(voices) == 7, f"expected 7 voices got {len(voices)}"
        keys = {v["key"] for v in voices}
        assert "match_my_voice" in keys and "literary" in keys

    def test_quota_snapshot_shape(self, user):
        r = requests.get(f"{API}/ai/agent/quota", headers=_headers(user["token"]), timeout=15)
        assert r.status_code == 200
        js = r.json()
        # fresh user should not be subscriber (unless mongo has state)
        if not js.get("unlimited"):
            assert set(js.keys()) >= {"unlimited", "used", "limit", "remaining"}
            assert js["limit"] == 5


class TestAgentQuotaDecrement:
    """Fresh user: 5 successful calls, then 6th returns 402 with agent_quota_exceeded."""

    def test_five_then_402(self):
        reg = _register()
        tok = reg["token"]
        h = _headers(tok)
        doc = requests.post(f"{API}/documents",
                            json={"title": "TEST_iter17_quota", "content": "<p>Some initial text.</p>"},
                            headers=h, timeout=15).json()
        did = doc["id"]

        remainings = []
        for i in range(5):
            r = requests.post(f"{API}/ai/agent/command",
                              json={"instruction": "make it shorter",
                                    "selected_text": f"call number {i} — this is a passage to rewrite.",
                                    "document_id": did},
                              headers=h, timeout=120)
            assert r.status_code == 200, f"call {i}: {r.status_code} {r.text[:150]}"
            js = r.json()
            assert js["result"]
            remainings.append(js["quota"]["remaining"])

        assert remainings == [4, 3, 2, 1, 0], f"expected 4->0, got {remainings}"

        # 6th call — MUST 402
        r6 = requests.post(f"{API}/ai/agent/command",
                           json={"instruction": "shorten",
                                 "selected_text": "sixth call, should be blocked",
                                 "document_id": did},
                           headers=h, timeout=60)
        assert r6.status_code == 402, f"expected 402, got {r6.status_code} {r6.text[:200]}"
        detail = r6.json().get("detail") or {}
        assert isinstance(detail, dict), f"detail should be dict, got {detail!r}"
        assert detail.get("code") == "agent_quota_exceeded"

        # And the quota snapshot shows 0 remaining
        q = requests.get(f"{API}/ai/agent/quota", headers=h, timeout=15).json()
        assert q.get("unlimited") is False
        assert q.get("remaining") == 0

    def test_subscriber_unlimited(self):
        """If a user has an active subscription doc, quota returns unlimited:true."""
        from pymongo import MongoClient
        from datetime import datetime, timezone, timedelta
        mongo_url = os.environ.get("MONGO_URL", "mongodb://localhost:27017")
        db_name = os.environ.get("DB_NAME", "test_database")

        reg = _register()
        tok = reg["token"]
        uid = reg["user"]["id"]

        client = MongoClient(mongo_url)
        db = client[db_name]
        db.subscriptions.insert_one({
            "id": str(uuid.uuid4()),
            "user_id": uid,
            "plan_id": "author_pro",
            "active": True,
            "pro_until": (datetime.now(timezone.utc) + timedelta(days=30)).isoformat(),
            "payments_count": 1,
            "created_at": datetime.now(timezone.utc).isoformat(),
            "updated_at": datetime.now(timezone.utc).isoformat(),
        })
        try:
            q = requests.get(f"{API}/ai/agent/quota", headers=_headers(tok), timeout=15).json()
            assert q.get("unlimited") is True, f"expected unlimited=True, got {q}"
        finally:
            db.subscriptions.delete_many({"user_id": uid})


class TestAgentChat:
    def test_chat_history_persists_and_clears(self, user):
        h = _headers(user["token"])
        did = user["doc_id"]
        r = requests.post(f"{API}/ai/agent/chat",
                          json={"document_id": did, "message": "In one sentence, describe this manuscript.",
                                "voice": "literary"},
                          headers=h, timeout=120)
        # If quota was exhausted by other tests, allow 402
        if r.status_code == 402:
            pytest.skip("quota already exhausted for this user (ok)")
        assert r.status_code == 200, r.text[:200]
        js = r.json()
        assert js["reply"]
        sid = js["session_id"]

        # history endpoint
        h_resp = requests.get(f"{API}/ai/agent/history/{did}", headers=h, timeout=15).json()
        assert len(h_resp["messages"]) >= 2
        # clear
        d = requests.delete(f"{API}/ai/agent/history/{did}", headers=h, timeout=15)
        assert d.status_code == 200
        h_resp2 = requests.get(f"{API}/ai/agent/history/{did}", headers=h, timeout=15).json()
        assert h_resp2["messages"] == []


# ---------- PER-CHAPTER AUDIOBOOK ----------

class TestChapterAudiobook:
    def test_preview_returns_chapters(self, user):
        h = _headers(user["token"])
        r = requests.get(f"{API}/documents/{user['doc_id']}/audiobook/chapters/preview", headers=h, timeout=15)
        assert r.status_code == 200
        js = r.json()
        assert js["total"] >= 2
        for ch in js["chapters"]:
            assert set(ch.keys()) >= {"index", "title", "char_count", "word_count"}

    def test_export_400_when_single_chapter(self, user):
        h = _headers(user["token"])
        # create a single-chapter doc
        d = requests.post(f"{API}/documents", json={"title": "TEST_iter17_single", "content": "<p>No headings here.</p>"}, headers=h, timeout=15).json()
        r = requests.post(f"{API}/documents/{d['id']}/audiobook/chapters", headers=h, timeout=30)
        assert r.status_code == 400
        assert "chapter" in r.text.lower()

    def test_export_400_when_empty(self, user):
        h = _headers(user["token"])
        d = requests.post(f"{API}/documents", json={"title": "TEST_iter17_empty", "content": ""}, headers=h, timeout=15).json()
        r = requests.post(f"{API}/documents/{d['id']}/audiobook/chapters", headers=h, timeout=30)
        assert r.status_code == 400


# ---------- AFFILIATE / REFERRALS / ADMIN GATES ----------

class TestAffiliateAndAdmin:
    def test_public_settings(self):
        r = requests.get(f"{API}/affiliate/settings", timeout=15)
        assert r.status_code == 200
        assert "signup_commission_percent" in r.json()

    def test_leaderboard(self):
        r = requests.get(f"{API}/referrals/leaderboard", timeout=15)
        assert r.status_code == 200
        assert "leaderboard" in r.json()

    def test_admin_settings_forbidden_for_regular(self, user):
        r = requests.put(f"{API}/admin/affiliate/settings", json={"enabled": True},
                         headers=_headers(user["token"]), timeout=15)
        assert r.status_code == 403

    def test_admin_settings_ok_for_admin(self, admin_token):
        r = requests.put(f"{API}/admin/affiliate/settings", json={"enabled": True},
                         headers=_headers(admin_token), timeout=15)
        assert r.status_code == 200

    def test_admin_commissions_gates(self, admin_token, user):
        # regular user 403
        r = requests.get(f"{API}/admin/commissions", headers=_headers(user["token"]), timeout=15)
        assert r.status_code == 403
        # admin 200
        r2 = requests.get(f"{API}/admin/commissions", headers=_headers(admin_token), timeout=15)
        assert r2.status_code == 200
        js = r2.json()
        assert "commissions" in js and "pending_by_user" in js

    def test_mark_paid_empty(self, admin_token):
        r = requests.post(f"{API}/admin/commissions/mark-paid",
                          json={"commission_ids": []}, headers=_headers(admin_token), timeout=15)
        assert r.status_code == 400

    def test_billing_commissions_shape(self, user):
        r = requests.get(f"{API}/billing/commissions", headers=_headers(user["token"]), timeout=15)
        assert r.status_code == 200
        js = r.json()
        assert "commissions" in js and "totals" in js
        assert set(js["totals"].keys()) >= {"pending", "paid", "count"}


# ---------- BILLING + LIVE STRIPE ----------

class TestBilling:
    def test_plans(self):
        r = requests.get(f"{API}/billing/plans", timeout=15)
        assert r.status_code == 200
        ids = {p["id"] for p in r.json()["plans"]}
        assert {"author_pro", "estate"}.issubset(ids)

    def test_diagnostics_live(self):
        r = requests.get(f"{API}/billing/diagnostics", timeout=15)
        assert r.status_code == 200
        js = r.json()
        assert js["subscription_billing_configured"] is True
        assert js["live_mode"] is True
        assert js["webhook_secret_configured"] is True

    def test_billing_me_default(self, user):
        r = requests.get(f"{API}/billing/me", headers=_headers(user["token"]), timeout=15)
        assert r.status_code == 200
        assert r.json()["active"] in (False, True)

    def test_live_checkout_creates_cs_live_session(self, user):
        h = _headers(user["token"])
        r = requests.post(f"{API}/billing/checkout",
                          json={"plan_id": "author_pro", "origin_url": BASE_URL},
                          headers=h, timeout=30)
        assert r.status_code == 200, r.text[:200]
        js = r.json()
        assert js["session_id"].startswith("cs_live_"), f"expected cs_live_, got {js['session_id']}"
        assert "stripe.com" in js["url"] or "checkout" in js["url"]

    def test_checkout_status_404_for_other_user(self, user):
        # create a session as user
        h = _headers(user["token"])
        r = requests.post(f"{API}/billing/checkout",
                          json={"plan_id": "author_pro", "origin_url": BASE_URL},
                          headers=h, timeout=30)
        sid = r.json()["session_id"]
        # try to fetch from a different user
        other = _register()
        r2 = requests.get(f"{API}/billing/checkout/status/{sid}",
                          headers=_headers(other["token"]), timeout=30)
        assert r2.status_code == 404

    def test_webhook_bad_signature(self):
        r = requests.post(f"{API}/webhook/stripe",
                          data=b'{"type":"checkout.session.completed"}',
                          headers={"Stripe-Signature": "t=0,v1=bogus", "Content-Type": "application/json"},
                          timeout=15)
        assert r.status_code == 400
        assert "signature" in r.text.lower() or "invalid" in r.text.lower()

    def test_portal_no_customer(self, user):
        r = requests.post(f"{API}/billing/portal", headers=_headers(user["token"]), timeout=15)
        assert r.status_code == 400
        assert "no stripe customer" in r.text.lower() or "customer" in r.text.lower()


# ---------- STRIPE CONNECT ----------

class TestStripeConnect:
    def test_connect_status_default_false(self, user):
        r = requests.get(f"{API}/affiliate/connect/status", headers=_headers(user["token"]), timeout=15)
        assert r.status_code == 200
        assert r.json().get("connected") is False

    def test_connect_onboard_503_or_actionable(self, user):
        r = requests.post(f"{API}/affiliate/connect/onboard",
                          json={"origin_url": BASE_URL},
                          headers=_headers(user["token"]), timeout=30)
        # Per spec: if Connect NOT enabled on platform account, expect 503 with the URL hint
        assert r.status_code in (503, 502), f"unexpected status {r.status_code}: {r.text[:200]}"
        if r.status_code == 503:
            body = r.text.lower()
            assert "connect" in body
            assert "dashboard.stripe.com" in body or "stripe" in body

    def test_auto_payout_empty_ids(self, admin_token):
        r = requests.post(f"{API}/admin/commissions/auto-payout",
                          json={"commission_ids": []}, headers=_headers(admin_token), timeout=15)
        assert r.status_code == 400

    def test_auto_payout_no_connect_account(self, admin_token, user):
        """Insert a pending commission for the fresh user (no Connect account),
        then request auto-payout — row should end up in `failed` with 'no_connect_account'."""
        from pymongo import MongoClient
        from datetime import datetime, timezone
        mongo_url = os.environ.get("MONGO_URL", "mongodb://localhost:27017")
        db_name = os.environ.get("DB_NAME", "test_database")
        db = MongoClient(mongo_url)[db_name]

        cid = str(uuid.uuid4())
        db.commissions.insert_one({
            "id": cid,
            "referrer_user_id": user["id"],   # user has no stripe_connect_account_id
            "referee_user_id": "someone",
            "referee_email": "x@y.z",
            "transaction_id": "TEST",
            "session_id": "TEST_iter17",
            "plan_id": "author_pro",
            "kind": "signup",
            "percent": 30.0,
            "amount": 5.70,
            "currency": "USD",
            "status": "pending",
            "created_at": datetime.now(timezone.utc).isoformat(),
        })
        try:
            r = requests.post(f"{API}/admin/commissions/auto-payout",
                              json={"commission_ids": [cid]},
                              headers=_headers(admin_token), timeout=30)
            assert r.status_code == 200, r.text[:200]
            js = r.json()
            assert js["paid"] == 0
            assert any(f["commission_id"] == cid and f["reason"] == "no_connect_account" for f in js["failed"])
        finally:
            db.commissions.delete_one({"id": cid})


# ---------- SUPER ADMIN GATES ----------

class TestAdminGates:
    """Every /api/admin/* endpoint must 403 for a regular user."""

    ENDPOINTS = [
        ("PUT",  "/admin/affiliate/settings", {"enabled": True}),
        ("GET",  "/admin/commissions", None),
        ("POST", "/admin/commissions/mark-paid", {"commission_ids": ["nope"]}),
        ("POST", "/admin/commissions/auto-payout", {"commission_ids": ["nope"]}),
    ]

    def test_all_admin_endpoints_403(self, user):
        h = _headers(user["token"])
        for method, path, body in self.ENDPOINTS:
            r = requests.request(method, f"{API}{path}", json=body, headers=h, timeout=15)
            assert r.status_code == 403, f"{method} {path} -> {r.status_code}"
